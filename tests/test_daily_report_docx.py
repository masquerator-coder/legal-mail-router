"""
日报 / 汇总报表 Word 附件 — 表格化生成测试

覆盖：
- _add_report_table 的表格结构（表头加粗、行列数、短行补空）
- _build_daily_report_docx 的五张表与空数据分支
- _build_target_summary_docx 的三张表与空数据分支
- send_report_mail 的 MIME 结构与临时文件清理
"""
# -*- coding: utf-8 -*-
import email
import os
import smtplib

import pytest
from docx import Document

from app.services.mail_forwarder import (
    _add_report_table,
    _build_daily_report_docx,
    _build_target_summary_docx,
    send_report_mail,
)


class TestAddReportTable:
    """通用表格助手"""

    def test_headers_and_rows(self):
        doc = Document()
        _add_report_table(doc, ["A", "B"], [["1", "2"], ["3", "4"]])
        t = doc.tables[0]
        assert len(t.columns) == 2
        assert len(t.rows) == 3          # 表头 + 2 行
        assert [c.text for c in t.rows[0].cells] == ["A", "B"]
        assert [c.text for c in t.rows[2].cells] == ["3", "4"]

    def test_header_runs_bold(self):
        doc = Document()
        _add_report_table(doc, ["A"], [])
        run = doc.tables[0].rows[0].cells[0].paragraphs[0].runs[0]
        assert run.bold is True

    def test_short_row_padded_and_long_row_truncated(self):
        """行长与表头不一致时不得抛异常"""
        doc = Document()
        _add_report_table(doc, ["A", "B", "C"], [["1"], ["1", "2", "3", "4"]])
        t = doc.tables[0]
        assert [c.text for c in t.rows[1].cells] == ["1", "", ""]
        assert [c.text for c in t.rows[2].cells] == ["1", "2", "3"]

    def test_empty_headers_no_table(self):
        doc = Document()
        assert _add_report_table(doc, [], [["1"]]) is None
        assert len(doc.tables) == 0


class TestDailyReportDocx:
    """管理员日报 Word 附件"""

    DATA = {
        "sys_name": "测试系统", "now_str": "2026年01月01日 09:00",
        "overview": [["处理总数", "12 封"], ["成功转发", "10 封"]],
        "cumulative": [["累计处理", "100 封"], ["累计成功率", "90.0%"]],
        "accounts": [["邮箱A", "a@x.com", "5 封", "01-01 08:00", "10 分钟"]],
        "system": [["调度引擎", "运行中"], ["服务端口", "8020"]],
        "errors": [["01-01 08:30", "某主题", "某原因"]],
    }

    def test_five_tables_with_expected_headers(self):
        path = _build_daily_report_docx(self.DATA)
        try:
            doc = Document(path)
            assert len(doc.tables) == 5
            assert [c.text for c in doc.tables[0].rows[0].cells] == ["指标", "数值"]
            assert [c.text for c in doc.tables[2].rows[0].cells] == \
                ["名称", "邮箱", "今日处理", "最近检查", "检查间隔"]
            assert [c.text for c in doc.tables[4].rows[0].cells] == ["时间", "主题", "原因"]
        finally:
            os.unlink(path)

    def test_title_and_headings(self):
        path = _build_daily_report_docx(self.DATA)
        try:
            doc = Document(path)
            text = "\n".join(p.text for p in doc.paragraphs)
            assert "测试系统 运行报告" in text
            assert "今日处理概况" in text and "系统状态" in text
        finally:
            os.unlink(path)

    def test_no_accounts_shows_hint(self):
        data = dict(self.DATA, accounts=[], errors=[])
        path = _build_daily_report_docx(data)
        try:
            doc = Document(path)
            text = "\n".join(p.text for p in doc.paragraphs)
            assert "（无启用的监控邮箱）" in text
            assert len(doc.tables) == 3      # 概况/累计/系统，无邮箱表与错误表
        finally:
            os.unlink(path)

    def test_account_row_cells(self):
        path = _build_daily_report_docx(self.DATA)
        try:
            doc = Document(path)
            assert [c.text for c in doc.tables[2].rows[1].cells] == \
                ["邮箱A", "a@x.com", "5 封", "01-01 08:00", "10 分钟"]
        finally:
            os.unlink(path)


class TestTargetSummaryDocx:
    """转发目标汇总 Word 附件"""

    DATA = {
        "sys_name": "测试系统", "now_str": "now", "day_str": "2026年01月01日",
        "target_email": "a@x.com", "name": "张律师", "total": 3, "failed": 0,
        "types": [["合同协议", "2 封", "66.7%"], ["政府信息公开", "1 封", "33.3%"]],
        "subjects": [["合同协议", "合同甲"]], "note": "",
    }

    def test_three_tables(self):
        path = _build_target_summary_docx(self.DATA)
        try:
            doc = Document(path)
            assert len(doc.tables) == 3
            assert [c.text for c in doc.tables[1].rows[0].cells] == ["文书类型", "数量", "占比"]
            assert [c.text for c in doc.tables[2].rows[0].cells] == ["文书类型", "邮件主题"]
        finally:
            os.unlink(path)

    def test_failed_row_included(self):
        path = _build_target_summary_docx(dict(self.DATA, failed=2))
        try:
            doc = Document(path)
            cells = [c.text for r in doc.tables[0].rows for c in r.cells]
            assert any("转发失败" == c for c in cells)
        finally:
            os.unlink(path)

    def test_empty_day_hint_and_no_list_table(self):
        data = dict(self.DATA, types=[], subjects=[], total=0)
        path = _build_target_summary_docx(data)
        try:
            doc = Document(path)
            text = "\n".join(p.text for p in doc.paragraphs)
            assert "（今日无转发记录）" in text
            assert len(doc.tables) == 1      # 只剩概况表
        finally:
            os.unlink(path)


class TestSendReportMail:
    """带 Word 附件的发送"""

    SMTP = {"host": "h", "port": 587, "username": "u", "password_encrypted": "enc"}

    @pytest.fixture
    def captured(self, monkeypatch):
        box = {}

        class FakeSMTP:
            def __init__(self, *a, **k):
                pass

            def starttls(self):
                pass

            def login(self, *a):
                pass

            def sendmail(self, frm, to, raw):
                box["raw"] = raw

            def quit(self):
                pass

        monkeypatch.setattr(smtplib, "SMTP", FakeSMTP)
        monkeypatch.setattr("app.config.decrypt", lambda s: "pw")
        return box

    def test_attaches_docx_and_sends_plain_body(self, captured):
        path = _build_target_summary_docx(TestTargetSummaryDocx.DATA)
        ok = send_report_mail(self.SMTP, "a@x.com", "主题", "正文摘要",
                              docx_path=path, attachment_name="汇总.docx")
        assert ok is True

        msg = email.message_from_string(captured["raw"])
        assert msg.get_content_type() == "multipart/mixed"
        # 非 ASCII 主题按 RFC 2047 编码，解码后应还原
        from email.header import decode_header, make_header
        assert str(make_header(decode_header(msg.get("Subject")))) == "主题"
        assert msg.get("To") == "a@x.com"
        # 正文 + 附件两部分
        assert msg.get_payload(decode=True) is None      # 容器
        plain = [p for p in msg.walk() if p.get_content_type() == "text/plain"]
        assert plain and plain[0].get_payload(decode=True).decode("utf-8") == "正文摘要"
        atts = [p for p in msg.walk() if p.get_filename()]
        assert len(atts) == 1
        assert atts[0].get_filename() == "汇总.docx"
        assert atts[0].get_payload(decode=True)[:2] == b"PK"   # 合法 docx(zip)

    def test_temp_file_removed_after_send(self, captured):
        path = _build_target_summary_docx(TestTargetSummaryDocx.DATA)
        assert os.path.exists(path)
        send_report_mail(self.SMTP, "a@x.com", "s", "b",
                         docx_path=path, attachment_name="x.docx")
        assert not os.path.exists(path)

    def test_without_docx_sends_simple_text(self, captured):
        ok = send_report_mail(self.SMTP, "a@x.com", "s", "纯文本")
        assert ok is True
        msg = email.message_from_string(captured["raw"])
        assert msg.get_content_type() == "text/plain"
        assert msg.get_payload(decode=True).decode("utf-8") == "纯文本"

    def test_temp_file_removed_even_on_failure(self, monkeypatch):
        """发送失败也必须清理临时附件，避免长期运行堆积"""
        path = _build_target_summary_docx(TestTargetSummaryDocx.DATA)

        class BoomSMTP:
            def __init__(self, *a, **k):
                raise OSError("connect failed")

        monkeypatch.setattr(smtplib, "SMTP", BoomSMTP)
        monkeypatch.setattr("app.config.decrypt", lambda s: "pw")
        ok = send_report_mail(self.SMTP, "a@x.com", "s", "b",
                              docx_path=path, attachment_name="x.docx")
        assert ok is False
        assert not os.path.exists(path)

    def test_ssl_port_uses_smtp_ssl(self, monkeypatch, captured):
        used = {}

        class FakeSSL:
            def __init__(self, host, port, timeout=None, context=None):
                used["port"] = port

            def login(self, *a):
                pass

            def sendmail(self, *a):
                pass

            def quit(self):
                pass

        monkeypatch.setattr(smtplib, "SMTP_SSL", FakeSSL)
        monkeypatch.setattr("app.config.decrypt", lambda s: "pw")
        ok = send_report_mail({**self.SMTP, "port": 465}, "a@x.com", "s", "b")
        assert ok is True
        assert used["port"] == 465