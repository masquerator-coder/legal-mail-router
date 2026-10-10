"""
转发目标总结邮件 — 按目标邮箱统计测试

覆盖：
- forward_records 逐目标记录的写入（成功/失败、多类型、姓名）
- _collect_target_summaries 按目标聚合：只统计当日、只统计成功、目标间不串
- 多类型邮件按类型分别计数
- 报表正文：类型分布、清单、隔离性（不含其他目标）
- 未启用时不发送
"""
# -*- coding: utf-8 -*-
from datetime import datetime, timedelta

import pytest

from app.services.scheduler import (
    _collect_target_summaries,
    _build_target_summary_body,
    _build_target_summary_brief,
    _record_forward_result,
    send_target_summary_reports,
)


def _disp_width(s: str) -> int:
    import unicodedata
    return sum(2 if unicodedata.east_asian_width(c) in ("W", "F") else 1 for c in s)


@pytest.fixture
def fwd_db():
    """独立内存库（不触碰真实 data 目录）"""
    from sqlalchemy import create_engine
    from sqlalchemy.orm import sessionmaker
    from app.database import Base
    from app.models import EmailLog, ForwardRecord, DefaultConfig  # noqa: F401  # 注册模型

    engine = create_engine("sqlite:///:memory:")
    Base.metadata.create_all(engine)
    db = sessionmaker(bind=engine)()
    yield db
    db.close()
    engine.dispose()


def _make_log(db, subject="测试邮件", doc_type="合同协议", doc_types=None):
    from app.models import EmailLog
    log = EmailLog(
        account_id=1, message_id=f"m-{subject}", subject=subject,
        doc_type=doc_type, doc_types=doc_types or doc_type, status="forwarded",
    )
    db.add(log)
    db.commit()
    return log


class TestRecordForwardResult:
    """逐目标转发记录写入"""

    def test_records_success(self, fwd_db):
        log = _make_log(fwd_db, "合同A", "合同协议")
        _record_forward_result(fwd_db, log, "a@x.com", "张律师", [{"doc_type": "合同协议"}], True)

        from app.models import ForwardRecord
        recs = fwd_db.query(ForwardRecord).all()
        assert len(recs) == 1
        assert recs[0].target_email == "a@x.com"
        assert recs[0].target_name == "张律师"
        assert recs[0].doc_type == "合同协议"
        assert recs[0].success is True
        assert recs[0].log_id == log.id

    def test_records_failure(self, fwd_db):
        log = _make_log(fwd_db, "合同B")
        _record_forward_result(fwd_db, log, "b@x.com", "", [{"doc_type": "合同协议"}],
                               False, "SMTP 认证失败")

        from app.models import ForwardRecord
        rec = fwd_db.query(ForwardRecord).one()
        assert rec.success is False

    def test_multiple_types_joined(self, fwd_db):
        log = _make_log(fwd_db, "多类型")
        _record_forward_result(
            fwd_db, log, "a@x.com", "张",
            [{"doc_type": "起诉状"}, {"doc_type": "证据材料"}], True,
        )
        from app.models import ForwardRecord
        rec = fwd_db.query(ForwardRecord).one()
        assert rec.doc_type == "起诉状"          # 首个作为主类型
        assert rec.doc_types == "起诉状,证据材料"

    def test_empty_analyses_falls_back_to_log(self, fwd_db):
        log = _make_log(fwd_db, "兜底", "律师函")
        _record_forward_result(fwd_db, log, "a@x.com", "", [], True)
        from app.models import ForwardRecord
        rec = fwd_db.query(ForwardRecord).one()
        assert rec.doc_type == "律师函"
        assert rec.doc_types == "律师函"

    def test_write_failure_does_not_raise(self, fwd_db):
        """写入异常必须被吞掉，不能影响转发主流程"""
        log = _make_log(fwd_db)
        bad_log = type("L", (), {"id": "not-an-int", "subject": "x",
                                 "doc_type": None, "doc_types": None})()
        _record_forward_result(fwd_db, bad_log, "a@x.com", "", [{"doc_type": "合同协议"}], True)
        # 未抛异常即通过


class TestCollectTargetSummaries:
    """按目标邮箱聚合"""

    def _seed(self, db):
        from app.models import ForwardRecord
        now = datetime.now()
        yesterday = now - timedelta(days=1)

        a1 = _make_log(db, "合同甲", "合同协议")
        a2 = _make_log(db, "公开答复", "政府信息公开")
        a3 = _make_log(db, "起诉+证据", "起诉状", doc_types="起诉状,证据材料")
        old = _make_log(db, "昨日合同", "合同协议")

        db.add_all([
            ForwardRecord(log_id=a1.id, target_email="a@x.com", target_name="张律师",
                          doc_type="合同协议", doc_types="合同协议", subject="合同甲",
                          success=True, created_at=now),
            ForwardRecord(log_id=a2.id, target_email="a@x.com", target_name="张律师",
                          doc_type="政府信息公开", doc_types="政府信息公开", subject="公开答复",
                          success=True, created_at=now),
            ForwardRecord(log_id=a3.id, target_email="a@x.com", target_name="张律师",
                          doc_type="起诉状", doc_types="起诉状,证据材料", subject="起诉+证据",
                          success=True, created_at=now),
            ForwardRecord(log_id=a1.id, target_email="b@x.com", target_name="李律师",
                          doc_type="合同协议", doc_types="合同协议", subject="合同甲",
                          success=True, created_at=now),
            ForwardRecord(log_id=a2.id, target_email="b@x.com", target_name="李律师",
                          doc_type="政府信息公开", doc_types="政府信息公开", subject="公开答复",
                          success=False, created_at=now),
            ForwardRecord(log_id=old.id, target_email="a@x.com", target_name="张律师",
                          doc_type="合同协议", doc_types="合同协议", subject="昨日合同",
                          success=True, created_at=yesterday),
        ])
        db.commit()

    def _today(self):
        return datetime.now().replace(hour=0, minute=0, second=0, microsecond=0)

    def test_only_today_counted(self, fwd_db):
        self._seed(fwd_db)
        s = _collect_target_summaries(fwd_db, self._today())
        assert s["a@x.com"]["total"] == 3  # 昨日那封不计入

    def test_failed_not_counted_in_total(self, fwd_db):
        self._seed(fwd_db)
        s = _collect_target_summaries(fwd_db, self._today())
        assert s["b@x.com"]["total"] == 1
        assert s["b@x.com"]["failed"] == 1

    def test_per_target_isolation(self, fwd_db):
        """每个目标只统计自己的，不串目标"""
        self._seed(fwd_db)
        s = _collect_target_summaries(fwd_db, self._today())
        assert s["b@x.com"]["types"] == {"合同协议": 1}
        assert "政府信息公开" not in s["b@x.com"]["types"]

    def test_multi_type_counts_each(self, fwd_db):
        self._seed(fwd_db)
        s = _collect_target_summaries(fwd_db, self._today())
        assert s["a@x.com"]["types"]["起诉状"] == 1
        assert s["a@x.com"]["types"]["证据材料"] == 1

    def test_name_captured(self, fwd_db):
        self._seed(fwd_db)
        s = _collect_target_summaries(fwd_db, self._today())
        assert s["a@x.com"]["name"] == "张律师"

    def test_subjects_listed(self, fwd_db):
        self._seed(fwd_db)
        s = _collect_target_summaries(fwd_db, self._today())
        subs = [subj for subj, _ in s["a@x.com"]["subjects"]]
        assert "合同甲" in subs and "公开答复" in subs
        assert "昨日合同" not in subs

    def test_empty_when_no_records(self, fwd_db):
        assert _collect_target_summaries(fwd_db, self._today()) == {}


class TestSummaryBody:
    """报表正文"""

    def _stat(self):
        return {
            "name": "张律师", "total": 3,
            "types": {"合同协议": 2, "政府信息公开": 1},
            "subjects": [("合同甲", "合同协议"), ("公开答复", "政府信息公开")],
            "failed": 0,
        }

    def test_contains_totals_and_types(self):
        body = _build_target_summary_body("系统", "a@x.com", "张律师", self._stat(),
                                          "2026年01月01日 18:00", "2026年01月01日")
        assert "转发总数：3 封" in body
        assert "合同协议" in body and "政府信息公开" in body
        assert "张律师 您好" in body
        assert "a@x.com" in body

    def test_type_bars_align_by_display_width(self):
        """类型分布列按显示宽度对齐（中文按 2 列）"""
        body = _build_target_summary_body("系统", "a@x.com", "张律师", self._stat(),
                                          "now", "today")
        bars = [ln for ln in body.split("\n") if ln.strip().endswith("%)")]
        assert len(bars) == 2
        offsets = {_disp_width(ln.split("封")[0].rstrip()) for ln in bars}
        assert len(offsets) == 1, f"类型列未对齐: {bars}"

    def test_no_other_target_leaked(self):
        body = _build_target_summary_body("系统", "a@x.com", "张律师", self._stat(),
                                          "now", "today")
        assert "b@x.com" not in body

    def test_failed_shown_when_present(self):
        stat = self._stat()
        stat["failed"] = 2
        body = _build_target_summary_body("系统", "a@x.com", "张律师", stat,
                                          "now", "today")
        assert "转发失败：2 封" in body

    def test_empty_day_message(self):
        stat = {"name": "", "total": 0, "types": {}, "subjects": [], "failed": 0}
        body = _build_target_summary_body("系统", "a@x.com", "", stat, "now", "today")
        assert "今日无转发记录" in body

    def test_long_subject_truncated(self):
        stat = self._stat()
        stat["subjects"] = [("超长主题" * 40, "合同协议")]
        body = _build_target_summary_body("系统", "a@x.com", "张", stat, "now", "today")
        for line in body.split("\n"):
            if line.strip().startswith("·"):
                assert len(line) < 100


class TestSendTargetSummaryReports:
    """发送入口"""

    def test_skips_when_disabled(self, fwd_db, monkeypatch):
        from app.models import DefaultConfig
        fwd_db.add(DefaultConfig(key="target_summary_enabled", value="false"))
        fwd_db.commit()

        called = []
        monkeypatch.setattr("app.database.SessionLocal", lambda: fwd_db)
        monkeypatch.setattr("app.services.mail_forwarder.send_report_mail",
                            lambda *a, **k: called.append(a) or True)

        send_target_summary_reports()
        assert called == []

    def test_skips_when_no_records(self, fwd_db, monkeypatch):
        from app.models import DefaultConfig
        fwd_db.add(DefaultConfig(key="target_summary_enabled", value="true"))
        fwd_db.commit()

        called = []
        monkeypatch.setattr("app.database.SessionLocal", lambda: fwd_db)
        monkeypatch.setattr("app.services.scheduler._get_smtp_config_for_report",
                            lambda db: {"host": "h", "port": 25, "username": "u",
                                        "password_encrypted": "p"})
        monkeypatch.setattr("app.services.mail_forwarder.send_report_mail",
                            lambda *a, **k: called.append(a) or True)

        send_target_summary_reports()
        assert called == []


class TestSummaryWordAttachment:
    """汇总邮件改为「摘要正文 + Word 表格附件」"""

    def _stat(self):
        return {
            "name": "张律师", "total": 3,
            "types": {"合同协议": 2, "政府信息公开": 1},
            "subjects": [("合同甲", "合同协议")],
            "failed": 0,
        }

    def test_brief_body_points_to_attachment(self):
        """摘要正文不应再罗列明细，只给数字与附件指引"""
        body = _build_target_summary_brief(
            "系统", "a@x.com", "张律师", self._stat(),
            "2026年01月01日 18:00", "2026年01月01日",
            attachment_name="系统 转发汇总 2026-01-01.docx",
        )
        assert "转发总数：3 封" in body
        assert "系统 转发汇总 2026-01-01.docx" in body
        assert "张律师 您好" in body
        # 明细列表标记不应出现在摘要里
        assert "· [合同协议] 合同甲" not in body

    def test_brief_body_no_attachment_name_falls_back(self):
        """未给附件名时不出现空的附件指引段"""
        body = _build_target_summary_brief(
            "系统", "a@x.com", "张律师", self._stat(),
            "now", "today", attachment_name="",
        )
        assert "附件《" not in body
        assert "转发总数：3 封" in body

    def test_brief_body_shows_failed(self):
        stat = self._stat()
        stat["failed"] = 2
        body = _build_target_summary_brief(
            "系统", "a@x.com", "张律师", stat, "now", "today", "x.docx")
        assert "转发失败：2 封" in body

    def test_docx_tables_built(self):
        """Word 附件应含概况/类型分布/转发清单三张表"""
        from docx import Document
        from app.services.mail_forwarder import _build_target_summary_docx
        import os

        path = _build_target_summary_docx({
            "sys_name": "系统", "now_str": "now", "day_str": "today",
            "target_email": "a@x.com", "name": "张律师", "total": 3, "failed": 0,
            "types": [["合同协议", "2 封", "66.7%"], ["政府信息公开", "1 封", "33.3%"]],
            "subjects": [["合同协议", "合同甲"]], "note": "",
        })
        try:
            doc = Document(path)
            assert len(doc.tables) == 3
            assert [c.text for c in doc.tables[1].rows[0].cells] == ["文书类型", "数量", "占比"]
            assert [c.text for c in doc.tables[1].rows[1].cells] == ["合同协议", "2 封", "66.7%"]
            # 只含本目标的邮箱，不得泄漏其他目标
            text = "\n".join(p.text for p in doc.paragraphs)
            assert "a@x.com" in text
            assert "b@x.com" not in text
        finally:
            os.unlink(path)

    def test_docx_empty_day_shows_hint(self):
        from docx import Document
        from app.services.mail_forwarder import _build_target_summary_docx
        import os

        path = _build_target_summary_docx({
            "sys_name": "系统", "now_str": "now", "day_str": "today",
            "target_email": "a@x.com", "name": "", "total": 0, "failed": 0,
            "types": [], "subjects": [], "note": "",
        })
        try:
            doc = Document(path)
            text = "\n".join(p.text for p in doc.paragraphs)
            assert "今日无转发记录" in text
        finally:
            os.unlink(path)